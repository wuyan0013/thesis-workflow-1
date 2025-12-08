# -*- coding: utf-8 -*-
"""
开题报告生成页面
上传论文Word文档，智能生成开题报告
"""

import sys
from pathlib import Path
import io
from datetime import date, timedelta

import streamlit as st
from docx import Document

sys.path.insert(0, str(Path(__file__).parent.parent))

from src.claude_client import ClaudeClient, ThesisGenerator
from src.docx_builder import DocxBuilder
from src.config import get_config, save_env_config

# ============ 页面配置 ============
st.set_page_config(
    page_title="开题报告 - ThesisAI",
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ============ 全局样式 - 统一黑金主题 ============
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

:root {
    --bg-primary: #0a0a0b;
    --bg-secondary: #111113;
    --bg-card: #16161a;
    --bg-card-hover: #1c1c21;
    --border-default: #27272a;
    --border-gold: #F59E0B;
    --text-primary: #FFFFFF;
    --text-secondary: #a1a1aa;
    --text-muted: #71717a;
    --gold-primary: #F59E0B;
    --gold-light: #FBBF24;
    --gold-dark: #D97706;
    --gradient-gold: linear-gradient(135deg, #F59E0B 0%, #D97706 50%, #F59E0B 100%);
    --gradient-gold-subtle: linear-gradient(135deg, rgba(245,158,11,0.15) 0%, rgba(217,119,6,0.1) 100%);
    --glow-gold: 0 0 30px rgba(245,158,11,0.3);
}

.stApp, html, body, [data-testid="stAppViewContainer"], .main > div {
    background: var(--bg-primary) !important;
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif !important;
}

.main > div, .block-container, [data-testid="stMainBlockContainer"] {
    max-width: 100% !important;
    width: 100% !important;
    padding: 2rem 4rem !important;
}

[data-testid="stHeader"], #MainMenu, footer, header, .stDeployButton {
    display: none !important;
}

section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #0d0d0f 0%, #111113 100%) !important;
    border-right: 1px solid var(--border-default) !important;
}

section[data-testid="stSidebar"] h1,
section[data-testid="stSidebar"] h2,
section[data-testid="stSidebar"] h3 {
    color: var(--gold-primary) !important;
    text-shadow: 0 0 20px rgba(245,158,11,0.3) !important;
}

h1, h2, h3, h4, h5, h6 {
    color: var(--text-primary) !important;
    font-family: 'Inter', sans-serif !important;
    font-weight: 700 !important;
}

p, span, label, div {
    color: var(--text-secondary) !important;
}

.stButton > button {
    background: var(--gradient-gold-subtle) !important;
    border: 1px solid rgba(245,158,11,0.3) !important;
    border-radius: 12px !important;
    color: var(--gold-primary) !important;
    font-weight: 600 !important;
    font-size: 0.95rem !important;
    padding: 0.875rem 1.5rem !important;
    min-height: 48px !important;
    transition: all 0.3s cubic-bezier(0.4, 0, 0.2, 1) !important;
}

.stButton > button:hover {
    background: linear-gradient(135deg, rgba(245,158,11,0.25) 0%, rgba(217,119,6,0.2) 100%) !important;
    border-color: var(--gold-primary) !important;
    box-shadow: var(--glow-gold) !important;
    transform: translateY(-2px) !important;
}

.stTextInput > div > div > input,
.stTextArea > div > div > textarea,
.stSelectbox > div > div > div {
    background: var(--bg-card) !important;
    border: 1px solid var(--border-default) !important;
    border-radius: 10px !important;
    color: var(--text-primary) !important;
}

.stTextInput > div > div > input:focus,
.stTextArea > div > div > textarea:focus {
    border-color: var(--gold-primary) !important;
    box-shadow: 0 0 0 3px rgba(245,158,11,0.15) !important;
}

div[data-testid="stExpander"] {
    background: var(--bg-card) !important;
    border: 1px solid var(--border-default) !important;
    border-radius: 16px !important;
}

div[data-testid="stExpander"]:hover {
    border-color: rgba(245,158,11,0.4) !important;
}

.stDownloadButton > button {
    background: var(--gradient-gold) !important;
    border: none !important;
    color: #000000 !important;
    font-weight: 700 !important;
    box-shadow: var(--glow-gold) !important;
}

.stDownloadButton > button:hover {
    box-shadow: 0 0 60px rgba(245,158,11,0.4) !important;
    transform: translateY(-3px) !important;
}

.uploadedFile {
    background: var(--bg-card) !important;
    border: 1px solid var(--border-default) !important;
    border-radius: 10px !important;
}

[data-testid="stFileUploader"] {
    background: var(--bg-card) !important;
    border: 2px dashed var(--border-default) !important;
    border-radius: 16px !important;
}

[data-testid="stFileUploader"]:hover {
    border-color: var(--gold-primary) !important;
}

.stSpinner > div {
    border-top-color: var(--gold-primary) !important;
}

.stProgress > div > div > div {
    background: var(--gradient-gold) !important;
}
</style>
""", unsafe_allow_html=True)


# ============ 状态初始化 ============
if "proposal_api_key" not in st.session_state:
    config = get_config()
    st.session_state.proposal_api_key = config.claude_api_key or ""
    st.session_state.proposal_api_base_url = config.claude_base_url or ""
    st.session_state.proposal_api_model = config.claude_model or "claude-opus-4-5-20251101"

if "parsed_thesis_proposal" not in st.session_state:
    st.session_state.parsed_thesis_proposal = None

if "proposal_result" not in st.session_state:
    st.session_state.proposal_result = None


def parse_docx(uploaded_file) -> dict:
    """解析Word文档，提取论文结构"""
    doc = Document(io.BytesIO(uploaded_file.read()))

    result = {
        "title": "",
        "abstract": "",
        "chapters": [],
        "full_text": ""
    }

    full_text_parts = []
    current_chapter = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        full_text_parts.append(text)
        style_name = para.style.name if para.style else ""

        if not result["title"] and ("Title" in style_name or len(full_text_parts) == 1):
            if len(text) < 50:
                result["title"] = text
                continue

        if "摘要" in text[:10] or "Abstract" in text[:10]:
            result["abstract"] = text
            continue

        is_heading = "Heading" in style_name
        is_chapter_marker = (
            text.startswith("一、") or text.startswith("二、") or text.startswith("三、") or
            text.startswith("四、") or text.startswith("五、") or text.startswith("六、") or
            text.startswith("第一章") or text.startswith("第二章") or text.startswith("第三章") or
            text.startswith("第四章") or text.startswith("第五章") or text.startswith("第六章")
        )

        if is_heading or is_chapter_marker:
            if current_chapter:
                result["chapters"].append(current_chapter)
            current_chapter = {"title": text, "content": ""}
        elif current_chapter:
            current_chapter["content"] += text + "\n"

    if current_chapter:
        result["chapters"].append(current_chapter)

    result["full_text"] = "\n".join(full_text_parts)

    if not result["title"]:
        result["title"] = "未检测到标题"

    return result


def create_proposal_docx(proposal_text: str, title: str) -> bytes:
    """创建开题报告Word文档"""
    doc = Document()
    doc.add_heading('开题报告', 0)
    doc.add_paragraph(f'论文题目：{title}')
    doc.add_paragraph('')

    for para in proposal_text.split('\n'):
        para = para.strip()
        if not para:
            continue
        if para.startswith('### '):
            doc.add_heading(para[4:], level=1)
        elif para.startswith('#### '):
            doc.add_heading(para[5:], level=2)
        elif para.startswith('**') and para.endswith('**'):
            p = doc.add_paragraph()
            run = p.add_run(para.strip('*'))
            run.bold = True
        elif para.startswith('- ') or para.startswith('* '):
            doc.add_paragraph(para[2:], style='List Bullet')
        else:
            doc.add_paragraph(para)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ============ 侧边栏 ============
with st.sidebar:
    st.markdown("## 📋 开题报告生成")
    st.caption("上传论文 → 智能生成开题报告")

    st.divider()

    with st.expander("⚙️ API 设置", expanded=not st.session_state.proposal_api_key):
        model_options = {
            "Claude Opus 4.5": "claude-opus-4-5-20251101",
            "Claude Sonnet 4": "claude-sonnet-4-20250514",
        }
        selected_model = st.selectbox(
            "模型",
            options=list(model_options.keys()),
            index=0 if st.session_state.proposal_api_model == "claude-opus-4-5-20251101" else 1,
            key="proposal_model_select"
        )
        st.session_state.proposal_api_model = model_options[selected_model]

        st.session_state.proposal_api_base_url = st.text_input(
            "API Base URL",
            value=st.session_state.proposal_api_base_url,
            placeholder="留空使用 .env 配置",
            key="proposal_base_url_input"
        )

        st.session_state.proposal_api_key = st.text_input(
            "API Key",
            value=st.session_state.proposal_api_key,
            type="password",
            placeholder="留空使用 .env 配置",
            key="proposal_api_key_input"
        )

        if st.button("保存到 .env", use_container_width=True, key="proposal_save_config"):
            if save_env_config(
                api_key=st.session_state.proposal_api_key,
                base_url=st.session_state.proposal_api_base_url,
                model=st.session_state.proposal_api_model
            ):
                st.success("已保存")
            else:
                st.error("保存失败")

    st.divider()

    st.markdown("### 专业信息")
    major = st.text_input("专业名称", placeholder="如：汉语言文学", key="proposal_major")

    st.markdown("### 研究进度时间")
    today = date.today()
    default_end = today + timedelta(days=150)

    col_date1, col_date2 = st.columns(2)
    with col_date1:
        start_date = st.date_input(
            "开始日期",
            value=today,
            min_value=date(2020, 1, 1),
            max_value=date(2030, 12, 31),
            key="proposal_start_date"
        )
    with col_date2:
        end_date = st.date_input(
            "结束日期",
            value=default_end,
            min_value=date(2020, 1, 1),
            max_value=date(2030, 12, 31),
            key="proposal_end_date"
        )

    st.divider()
    st.caption("生成内容仅供参考")


# ============ 主内容区 ============
st.markdown("""
<div style="margin-bottom: 2rem;">
    <h1 style="font-size: 2.5rem; margin-bottom: 0.5rem;">📋 开题报告生成器</h1>
    <p style="color: #A1A1AA; font-size: 1.1rem;">上传您的论文Word文档，AI智能生成开题报告</p>
</div>
""", unsafe_allow_html=True)

st.markdown("""
<div style="background: rgba(24, 24, 27, 0.6); border: 1px solid rgba(63, 63, 70, 0.4); border-radius: 16px; padding: 1.5rem; margin-bottom: 2rem;">
    <div style="color: #F59E0B; font-size: 1rem; font-weight: 600; margin-bottom: 1rem;">📋 开题报告结构（广西开放大学规范）</div>
    <div style="display: grid; grid-template-columns: repeat(5, 1fr); gap: 1rem;">
        <div style="background: rgba(10, 10, 12, 0.8); border-radius: 10px; padding: 1rem;">
            <div style="color: #FFFFFF; font-weight: 600; margin-bottom: 0.5rem;">一、选题目的</div>
            <div style="color: #71717A; font-size: 0.85rem;">背景+目的+意义</div>
        </div>
        <div style="background: rgba(10, 10, 12, 0.8); border-radius: 10px; padding: 1rem;">
            <div style="color: #FFFFFF; font-weight: 600; margin-bottom: 0.5rem;">二、研究内容</div>
            <div style="color: #71717A; font-size: 0.85rem;">内容+思路+方法</div>
        </div>
        <div style="background: rgba(10, 10, 12, 0.8); border-radius: 10px; padding: 1rem;">
            <div style="color: #FFFFFF; font-weight: 600; margin-bottom: 0.5rem;">三、进度安排</div>
            <div style="color: #71717A; font-size: 0.85rem;">5阶段时间规划</div>
        </div>
        <div style="background: rgba(10, 10, 12, 0.8); border-radius: 10px; padding: 1rem;">
            <div style="color: #FFFFFF; font-weight: 600; margin-bottom: 0.5rem;">四、论文框架</div>
            <div style="color: #71717A; font-size: 0.85rem;">目录结构预览</div>
        </div>
        <div style="background: rgba(10, 10, 12, 0.8); border-radius: 10px; padding: 1rem;">
            <div style="color: #FFFFFF; font-weight: 600; margin-bottom: 0.5rem;">五、参考文献</div>
            <div style="color: #71717A; font-size: 0.85rem;">GB/T 7714格式</div>
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

col_left, col_right = st.columns([1, 1], gap="large")

with col_left:
    st.markdown("### 上传论文")

    uploaded_file = st.file_uploader(
        "拖拽或点击上传 Word 文档",
        type=["docx"],
        key="proposal_thesis_uploader",
        help="支持 .docx 格式的论文文档"
    )

    if uploaded_file:
        with st.spinner("正在解析文档..."):
            st.session_state.parsed_thesis_proposal = parse_docx(uploaded_file)

        thesis = st.session_state.parsed_thesis_proposal

        st.success("文档解析完成")

        st.markdown("#### 文档信息")
        col1, col2 = st.columns(2)
        with col1:
            st.metric("论文标题", thesis["title"][:20] + "..." if len(thesis["title"]) > 20 else thesis["title"])
        with col2:
            st.metric("检测章节", f"{len(thesis['chapters'])} 章")

        if thesis["chapters"]:
            with st.expander("章节结构", expanded=False):
                for i, ch in enumerate(thesis["chapters"], 1):
                    ch_title = ch["title"][:30] + "..." if len(ch["title"]) > 30 else ch["title"]
                    ch_len = len(ch["content"])
                    st.caption(f"{i}. {ch_title} ({ch_len} 字)")

        st.markdown("###")
        if st.button("生成开题报告", type="primary", use_container_width=True, key="generate_proposal"):
            try:
                with st.spinner("正在生成开题报告，请稍候..."):
                    client = ClaudeClient(
                        use_cli=True,
                        api_key=st.session_state.proposal_api_key,
                        base_url=st.session_state.proposal_api_base_url,
                        model=st.session_state.proposal_api_model
                    )
                    generator = ThesisGenerator(client)

                    result = generator.generate_proposal_report(
                        thesis_content=thesis["full_text"],
                        title=thesis["title"],
                        major=major if major else "",
                        start_date=start_date.strftime("%Y年%m月%d日"),
                        end_date=end_date.strftime("%Y年%m月%d日")
                    )

                    st.session_state.proposal_result = result["content"]
                    st.success("开题报告生成完成！")
            except Exception as e:
                st.error(f"生成失败: {e}")

with col_right:
    st.markdown("### 生成结果")

    if st.session_state.proposal_result:
        result_text = st.session_state.proposal_result
        thesis = st.session_state.parsed_thesis_proposal

        st.text_area(
            "开题报告内容",
            value=result_text,
            height=500,
            key="proposal_result_text"
        )

        col_dl1, col_dl2 = st.columns(2)
        with col_dl1:
            docx_bytes = create_proposal_docx(result_text, thesis["title"])
            st.download_button(
                label="下载 Word 文档",
                data=docx_bytes,
                file_name=f"开题报告_{thesis['title'][:20]}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
        with col_dl2:
            st.download_button(
                label="下载 Markdown",
                data=result_text,
                file_name=f"开题报告_{thesis['title'][:20]}.md",
                mime="text/markdown",
                use_container_width=True
            )
    else:
        st.markdown("""
        <div style="background: var(--bg-card); border: 2px dashed var(--border-default); border-radius: 16px; padding: 4rem 2rem; text-align: center;">
            <div style="font-size: 3rem; margin-bottom: 1rem;">📋</div>
            <div style="color: var(--text-muted); font-size: 1rem;">上传论文并点击生成按钮</div>
            <div style="color: var(--text-muted); font-size: 0.85rem; margin-top: 0.5rem;">开题报告将在这里显示</div>
        </div>
        """, unsafe_allow_html=True)
