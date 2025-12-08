# -*- coding: utf-8 -*-
"""
Word文档生成器
适配国家开放大学本科毕业论文格式规范
"""

from pathlib import Path
from typing import Optional, List
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

from .config import get_config


class DocxBuilder:
    """
    Word文档构建器
    适配国家开放大学本科毕业论文格式
    """

    # 中文数字映射
    CN_NUMBERS = ['一', '二', '三', '四', '五', '六', '七', '八', '九', '十',
                  '十一', '十二', '十三', '十四', '十五', '十六', '十七', '十八', '十九', '二十']

    # 中文括号数字映射
    CN_BRACKETS = ['（一）', '（二）', '（三）', '（四）', '（五）', '（六）', '（七）', '（八）', '（九）', '（十）',
                   '（十一）', '（十二）', '（十三）', '（十四）', '（十五）']

    def __init__(self, template_path: Optional[str] = None):
        """
        初始化文档构建器

        Args:
            template_path: Word模板文件路径，如果为None则创建新文档
        """
        self.config = get_config()

        if template_path and Path(template_path).exists():
            self.doc = Document(template_path)
        else:
            self.doc = Document()
            self._setup_document()

        # 章节计数器
        self._heading1_count = 0
        self._heading2_count = 0
        self._heading3_count = 0
        self._heading4_count = 0

    def _setup_document(self) -> None:
        """设置文档基本格式（国开大规范）"""
        # 页面设置 - A4纸
        section = self.doc.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)

        # 页边距（按规范要求：上下2.6cm，左右3cm，页眉页脚1.8cm）
        section.top_margin = Cm(2.6)
        section.bottom_margin = Cm(2.6)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(3.0)
        section.header_distance = Cm(1.8)
        section.footer_distance = Cm(1.8)

        # 设置样式
        self._setup_styles()

    def _setup_styles(self) -> None:
        """设置文档样式（按完整格式要求）"""
        styles = self.doc.styles

        # 正文样式 - 小四号宋体(12pt)，1.5倍行距，段间距0
        normal_style = styles['Normal']
        normal_style.font.name = 'Times New Roman'
        normal_style.font.size = Pt(12)
        normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        normal_style.paragraph_format.line_spacing = 1.5
        normal_style.paragraph_format.space_before = Pt(0)
        normal_style.paragraph_format.space_after = Pt(0)

        # 一级标题 - 小二号(18pt)黑体加粗，单倍行距，段前段后各24磅
        h1_style = styles['Heading 1']
        h1_style.font.name = 'Times New Roman'
        h1_style.font.size = Pt(18)
        h1_style.font.bold = True
        h1_style._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        h1_style.paragraph_format.space_before = Pt(24)
        h1_style.paragraph_format.space_after = Pt(24)
        h1_style.paragraph_format.line_spacing = 1.0

        # 二级标题 - 小三号(15pt)黑体加粗，单倍行距，段前段后各18磅
        h2_style = styles['Heading 2']
        h2_style.font.name = 'Times New Roman'
        h2_style.font.size = Pt(15)
        h2_style.font.bold = True
        h2_style._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        h2_style.paragraph_format.space_before = Pt(18)
        h2_style.paragraph_format.space_after = Pt(18)
        h2_style.paragraph_format.line_spacing = 1.0

        # 三级标题 - 小四号(12pt)黑体加粗，单倍行距，段前段后各12磅
        h3_style = styles['Heading 3']
        h3_style.font.name = 'Times New Roman'
        h3_style.font.size = Pt(12)
        h3_style.font.bold = True
        h3_style._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        h3_style.paragraph_format.space_before = Pt(12)
        h3_style.paragraph_format.space_after = Pt(12)
        h3_style.paragraph_format.line_spacing = 1.0

    def _set_run_font(self, run, font_name: str, east_asia_font: str, size_pt: float, bold: bool = False):
        """设置run的字体 - 完整设置所有字体属性确保中文字体正确显示"""
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        # 必须设置全部四个字体属性，否则Word会使用回退字体
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:eastAsia'), east_asia_font)
        rFonts.set(qn('w:cs'), east_asia_font)

    def add_toc(self, title: str = "目录") -> None:
        """
        添加目录（Word自动目录）

        Args:
            title: 目录标题
        """
        # 目录标题 - 小二号(18pt)黑体加粗，居中，单倍行距，段前段后各24磅
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(24)
        # 字间加空格
        spaced_title = " ".join(title)  # "目 录"
        run = p.add_run(spaced_title)
        self._set_run_font(run, 'Times New Roman', '黑体', 18, bold=True)

        # 插入目录域代码
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')

        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = 'TOC \\o "1-3" \\h \\z \\u'

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')

        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')

        run._element.append(fldChar1)
        run._element.append(instrText)
        run._element.append(fldChar2)
        run._element.append(fldChar3)

        # 添加提示文字
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run('（请在Word中右键点击目录，选择"更新域"以生成目录）')
        self._set_run_font(run, 'Times New Roman', '宋体', 10.5)

        self.doc.add_page_break()

    def add_abstract(
        self,
        abstract_cn: str,
        keywords_cn: List[str],
        abstract_en: str = "",
        keywords_en: List[str] = None,
    ) -> None:
        """
        添加摘要

        Args:
            abstract_cn: 中文摘要
            keywords_cn: 中文关键词列表
            abstract_en: 英文摘要（可选）
            keywords_en: 英文关键词列表（可选）
        """
        # 摘要标题 - 小二号(18pt)黑体加粗，居中，单倍行距，段前段后各24磅
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(24)
        run = p.add_run("摘 要")  # 字间空一格
        self._set_run_font(run, 'Times New Roman', '黑体', 18, bold=True)

        # 摘要内容 - 四号宋体(14pt)，1.5倍行距，支持多段落
        # 按换行符分割摘要内容，每段都要首行缩进
        paragraphs = abstract_cn.strip().split('\n')
        for para_text in paragraphs:
            para_text = para_text.strip()
            if para_text:
                p = self.doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0.74)  # 首行缩进两字符
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(para_text)
                self._set_run_font(run, 'Times New Roman', '宋体', 14)

        # 空一行
        self.doc.add_paragraph()

        # 关键词 - "关键词"三字四号黑体，内容四号宋体，首行缩进两格
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)  # 首行缩进
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run("关键词：")
        self._set_run_font(run, 'Times New Roman', '黑体', 14, bold=False)
        run = p.add_run("；".join(keywords_cn))
        self._set_run_font(run, 'Times New Roman', '宋体', 14)

        self.doc.add_page_break()

        # 英文摘要（如果有）
        if abstract_en:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("Abstract")
            self._set_run_font(run, 'Times New Roman', 'Times New Roman', 18, bold=True)

            self.doc.add_paragraph()

            p = self.doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.74)
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run(abstract_en)
            self._set_run_font(run, 'Times New Roman', 'Times New Roman', 14)

            self.doc.add_paragraph()

            if keywords_en:
                p = self.doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.5
                run = p.add_run("Keywords: ")
                run.font.bold = True
                run = p.add_run("; ".join(keywords_en))

            self.doc.add_page_break()

    def add_heading1(self, title: str, auto_number: bool = True) -> None:
        """
        添加一级标题（一、二、三、格式）- 小二号(18pt)黑体加粗

        Args:
            title: 标题文本（不含编号）
            auto_number: 是否自动编号
        """
        if auto_number:
            self._heading1_count += 1
            self._heading2_count = 0  # 重置二级计数
            self._heading3_count = 0  # 重置三级计数
            number = self.CN_NUMBERS[self._heading1_count - 1]
            full_title = f"{number}、{title}"
        else:
            full_title = title

        # 不使用样式，完全手动设置格式，避免样式覆盖字体
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(24)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(full_title)
        self._set_run_font(run, 'Times New Roman', '黑体', 18, bold=True)

    def add_heading2(self, title: str, auto_number: bool = True) -> None:
        """
        添加二级标题（（一）（二）（三）格式）- 小三号(15pt)黑体加粗

        Args:
            title: 标题文本（不含编号）
            auto_number: 是否自动编号
        """
        if auto_number:
            self._heading2_count += 1
            self._heading3_count = 0  # 重置三级计数
            number = self.CN_BRACKETS[self._heading2_count - 1]
            full_title = f"{number}{title}"
        else:
            full_title = title

        # 不使用样式，完全手动设置格式，避免样式覆盖字体
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(18)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(full_title)
        self._set_run_font(run, 'Times New Roman', '黑体', 15, bold=True)

    def add_heading3(self, title: str, auto_number: bool = True) -> None:
        """
        添加三级标题（1、2、3格式）- 小四号(12pt)黑体加粗

        Args:
            title: 标题文本（不含编号）
            auto_number: 是否自动编号
        """
        if auto_number:
            self._heading3_count += 1
            self._heading4_count = 0  # 重置四级计数
            full_title = f"{self._heading3_count}、{title}"
        else:
            full_title = title

        # 不使用样式，完全手动设置格式，避免样式覆盖字体
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(full_title)
        self._set_run_font(run, 'Times New Roman', '黑体', 12, bold=True)

    def add_heading4(self, title: str, auto_number: bool = True) -> None:
        """
        添加四级标题（（1）（2）格式）- 小四号楷体加粗，行前空两格

        Args:
            title: 标题文本（不含编号）
            auto_number: 是否自动编号
        """
        if auto_number:
            self._heading4_count += 1
            full_title = f"（{self._heading4_count}）{title}"
        else:
            full_title = title

        # 四级标题不使用 Heading 样式，手动设置格式
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)  # 行前空两格
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(full_title)
        self._set_run_font(run, 'Times New Roman', '楷体', 12, bold=True)

    def add_paragraph(self, text: str, first_line_indent: bool = True) -> None:
        """
        添加正文段落

        Args:
            text: 段落文本
            first_line_indent: 是否首行缩进
        """
        p = self.doc.add_paragraph()
        if first_line_indent:
            p.paragraph_format.first_line_indent = Cm(0.74)  # 两字符缩进
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        run = p.add_run(text)
        self._set_run_font(run, 'Times New Roman', '宋体', 12)

    def add_content(self, content: str) -> None:
        """
        添加内容（自动解析Markdown格式的标题和段落）

        Args:
            content: Markdown格式的内容
        """
        lines = content.strip().split('\n')
        current_paragraph = []

        for line in lines:
            line = line.strip()

            # 空行表示段落结束
            if not line:
                if current_paragraph:
                    self.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                continue

            # 解析标题
            if line.startswith('### '):
                if current_paragraph:
                    self.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                self.add_heading3(line[4:])
            elif line.startswith('## '):
                if current_paragraph:
                    self.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                self.add_heading2(line[3:])
            elif line.startswith('# '):
                if current_paragraph:
                    self.add_paragraph(' '.join(current_paragraph))
                    current_paragraph = []
                self.add_heading1(line[2:])
            else:
                # 普通文本，累积到当前段落
                current_paragraph.append(line)

        # 处理最后一个段落
        if current_paragraph:
            self.add_paragraph(' '.join(current_paragraph))

    def add_references(self, references: List[str], title: str = "参考文献") -> None:
        """
        添加参考文献（单独成页）

        Args:
            references: 参考文献列表
            title: 标题（默认"参考文献"）
        """
        # 分页符确保参考文献单独成页
        self.doc.add_page_break()

        # 参考文献标题 - 18pt黑体，居中
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title)
        self._set_run_font(run, 'Times New Roman', '黑体', 18, bold=False)

        # 参考文献内容 - 12pt宋体
        for i, ref in enumerate(references, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            # 如果引用已经有编号，直接使用；否则添加编号
            if not ref.strip().startswith('['):
                ref_text = f"[{i}]{ref}"
            else:
                ref_text = ref

            run = p.add_run(ref_text)
            self._set_run_font(run, 'Times New Roman', '宋体', 12)

    def add_appendix(self, title: str, content: str) -> None:
        """
        添加附录

        Args:
            title: 附录标题
            content: 附录内容
        """
        self.doc.add_page_break()

        # 附录标题
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        self._set_run_font(run, 'Times New Roman', '黑体', 18)

        self.doc.add_paragraph()

        # 附录内容
        for para_text in content.split('\n\n'):
            if para_text.strip():
                self.add_paragraph(para_text.strip())

    def add_page_break(self) -> None:
        """添加分页符"""
        self.doc.add_page_break()

    def save(self, filename: Optional[str] = None) -> Path:
        """
        保存文档

        Args:
            filename: 文件名（不含扩展名）

        Returns:
            保存的文件路径
        """
        output_dir = self.config.output_dir
        output_dir.mkdir(parents=True, exist_ok=True)

        if filename is None:
            filename = f"thesis_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

        # 清理文件名中的非法字符
        filename = "".join(c for c in filename if c not in r'\/:*?"<>|')

        output_path = output_dir / f"{filename}.docx"
        self.doc.save(str(output_path))

        return output_path
    def build_proposal_report(self, content: str, title: str = "") -> Path:
        """
        构建开题报告文档

        Args:
            content: Markdown格式的开题报告内容
            title: 论文标题

        Returns:
            保存的文件路径
        """
        lines = content.split('\n')

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # 一级标题（### 一、...）
            if line.startswith('### 一、') or line.startswith('### 二、') or line.startswith('### 三、') or line.startswith('### 四、') or line.startswith('### 五、'):
                heading_text = line.replace('### ', '').strip()
                self.add_heading1(heading_text, auto_number=False)

            # 二级标题（#### 1. ...）
            elif line.startswith('#### '):
                heading_text = line.replace('#### ', '').strip()
                self.add_heading2(heading_text, auto_number=False)

            # 三级标题（**...：**）
            elif line.startswith('**') and line.endswith('**') and '：' in line:
                heading_text = line.strip('*').strip()
                para = self.doc.add_paragraph()
                run = para.add_run(heading_text)
                run.bold = True
                self._set_run_font(run, 'Times New Roman', '黑体', 12)

            # 列表项
            elif line.startswith('- ') or line.startswith('* '):
                item_text = line[2:].strip()
                self.add_paragraph('• ' + item_text, first_line_indent=False)

            # 普通段落
            elif not line.startswith('#'):
                self.add_paragraph(line)

        # 保存文档
        filename = f"开题报告_{title}" if title else f"开题报告_{self.config.output_dir.name}"
        return self.save(filename)
